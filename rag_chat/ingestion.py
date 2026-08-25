"""Read files, split text into chunks, and add chunks to ChromaDB."""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from typing import BinaryIO

from rag_chat.client import get_embeddings
from rag_chat.config import get_settings
from rag_chat.store import get_collection

SUPPORTED_SUFFIXES = {".txt", ".pdf", ".docx"}


@dataclass(frozen=True)
class IndexedDocument:
    """A short summary of one indexed document."""

    source: str
    chunks: int
    added: int


def split_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> list[str]:
    """Split text into overlapping character chunks suitable for retrieval."""
    if chunk_size <= chunk_overlap:
        raise ValueError("chunk_size must be larger than chunk_overlap.")
    return [
        text[start : start + chunk_size]
        for start in range(0, len(text), chunk_size - chunk_overlap)
        if text[start : start + chunk_size].strip()
    ]


def extract_text(file: BinaryIO, filename: str) -> str:
    """Return readable text from a TXT, PDF, or DOCX file object."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        text = file.read().decode("utf-8", errors="replace")
    elif suffix == ".pdf":
        from pypdf import PdfReader

        text = "\n".join(page.extract_text() or "" for page in PdfReader(file).pages)
    elif suffix == ".docx":
        import docx

        text = "\n".join(paragraph.text for paragraph in docx.Document(file).paragraphs)
    else:
        raise ValueError("Supported file types are TXT, PDF, and DOCX.")
    if not text.strip():
        raise ValueError("No extractable text was found in this file.")
    return text


def index_text(text: str, source: str) -> IndexedDocument:
    """Embed and persist new chunks; re-uploading identical content is safe."""
    settings = get_settings()
    chunks = split_text(text, settings.chunk_size, settings.chunk_overlap)
    document_id = sha256(text.encode("utf-8")).hexdigest()
    ids = [f"{document_id}_chunk_{index}" for index in range(len(chunks))]
    collection = get_collection()
    existing = set(collection.get(ids=ids)["ids"])
    new_chunks = [(chunk_id, chunk) for chunk_id, chunk in zip(ids, chunks) if chunk_id not in existing]
    if new_chunks:
        collection.upsert(
            ids=[chunk_id for chunk_id, _ in new_chunks],
            documents=[chunk for _, chunk in new_chunks],
            embeddings=get_embeddings([chunk for _, chunk in new_chunks]),
            metadatas=[{"source": source, "document_id": document_id}] * len(new_chunks),
        )
    return IndexedDocument(source=source, chunks=len(chunks), added=len(new_chunks))


def index_path(path: Path) -> IndexedDocument:
    """Read and index one local document path."""
    with path.open("rb") as file:
        return index_text(extract_text(file, path.name), path.name)


def index_directory(directory: Path) -> list[IndexedDocument]:
    """Index every supported file in a directory, in filename order."""
    if not directory.exists():
        return []
    return [index_path(path) for path in sorted(directory.iterdir()) if path.suffix.lower() in SUPPORTED_SUFFIXES]